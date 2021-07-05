from docx import Document
import xlwt
from xlwt import *
import os
import docx

col_spec=0
col_section=col_spec+1
col_cal=col_section+1
word_root="/home/amai/Work/MY-github/work_tool/python/wordhelper"

def parseword(file_dir, sheet, index):
    doc = Document(file_dir)
    current_sec=""
    current_cal=""

    # for paragraph in iter_paragraphs(doc):
    #     try:
    #         num_pr = paragraph._p.pPr.numPr
    #         if num_pr is not None:
    #             print(num_pr.__dict__) 
    #             print(paragraph.text)
    #     except AttributeError:
    #         continue

    for p in doc.paragraphs:  
        style_name = p.style.name
        if style_name.startswith('Heading'):
            if(current_cal != "") :
                print(current_sec)
                print(current_cal)
                sheet.write(index, col_section, current_sec)
                sheet.write(index, col_cal, current_cal)
                index=index+1
                current_cal=""
            current_sec=p.text
        if("标定" in p.text):
            current_cal = current_cal + p.text
    return index

# def iter_paragraphs(parent, recursive=True):
#     """
#     Yield each paragraph and table child within *parent*, in document order.
#     Each returned value is an instance of Paragraph. *parent*
#     would most commonly be a reference to a main Document object, but
#     also works for a _Cell object, which itself can contain paragraphs and tables.
#     """
#     if isinstance(parent, docx.document.Document):
#         parent_elm = parent.element.body
#     elif isinstance(parent, docx.table._Cell):
#         parent_elm = parent._tc
#     else:
#         raise TypeError(repr(type(parent)))

#     for child in parent_elm.iterchildren():
#         if isinstance(child, docx.oxml.text.paragraph.CT_P):
#             yield docx.text.paragraph.Paragraph(child, parent)
#         elif isinstance(child, docx.oxml.table.CT_Tbl):
#             if recursive:
#                 try :
#                     table = docx.table.Table(child, parent)
#                     for row in table.rows:
#                         for cell in row.cells:
#                             for child_paragraph in iter_paragraphs(cell):
#                                 yield child_paragraph
#                 except AttributeError:
#                     continue

def main():
    index=0
    cal_book = xlwt.Workbook(encoding = 'utf-8')
    cal_sheet = cal_book.add_sheet('cal_sheet')
    filelist=os.listdir(word_root)

    for e in filelist:
        if ".docx" in e :
            # print(e)
            cal_sheet.write(index, col_spec, e.replace(".docx", ""))
            index=parseword(word_root+"/"+e, cal_sheet, index)+1

    cal_book.save('CalId.xls')

main()